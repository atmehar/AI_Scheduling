import os
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent.parent.parent

test_cases_data = [
    {
        "id": "TC-POS-01",
        "feature": "Slot Discovery",
        "preconditions": "Provider active, working hours set Mon-Fri 09:00-17:00, 30m service + 10m buffer.",
        "steps": "1. GET /api/providers/{id}/slots/ with from_date, to_date, service_id, and tz=America/New_York",
        "input": "service_id=1, from=2026-11-02, to=2026-11-02, tz=America/New_York",
        "expected": "HTTP 200 OK. Dynamic slots returned in customer timezone with 40-minute increments.",
        "actual": "HTTP 200 OK. 12 valid slots returned with correct display_start and UTC start_at.",
        "status": "PASS",
        "defect_id": "None",
        "evidence_link": "apps.scheduling.tests.SchedulingAndBookingEngineTests.test_dynamic_slot_generation"
    },
    {
        "id": "TC-POS-02",
        "feature": "Booking Creation",
        "preconditions": "Customer authenticated, target slot open.",
        "steps": "1. POST /api/bookings/ with provider_id, service_id, start_at, and UUID Idempotency-Key header.",
        "input": "provider_id=1, service_id=1, start_at='2026-11-02T14:00:00Z', Idempotency-Key='UUIDv4'",
        "expected": "HTTP 201 Created. Booking status 'confirmed', BookingEvent 'created' logged, 24h/1h reminders queued.",
        "actual": "HTTP 201 Created. Booking ID generated, audit event created, 2 reminders scheduled.",
        "status": "PASS",
        "defect_id": "None",
        "evidence_link": "apps.scheduling.test_qa_suite.FullQASuiteTests.test_positive_full_journey"
    },
    {
        "id": "TC-POS-03",
        "feature": "Atomic Reschedule",
        "preconditions": "Booking confirmed, customer authenticated.",
        "steps": "1. PATCH /api/bookings/{id}/reschedule/ with valid new_start_at.",
        "input": "new_start_at='2026-11-02T15:20:00Z'",
        "expected": "HTTP 200 OK. start_at and end_at shifted, reschedule_count incremented, reminders rescheduled.",
        "actual": "HTTP 200 OK. reschedule_count=1, reminders updated to new slot time.",
        "status": "PASS",
        "defect_id": "None",
        "evidence_link": "apps.scheduling.test_qa_suite.FullQASuiteTests.test_positive_full_journey"
    },
    {
        "id": "TC-POS-04",
        "feature": "Complete Appointment",
        "preconditions": "Booking confirmed, provider authenticated.",
        "steps": "1. POST /api/bookings/{id}/complete/ by assigned provider.",
        "input": "booking_id=1, provider_user=Doctor",
        "expected": "HTTP 200 OK. Booking status transitions to 'completed'.",
        "actual": "HTTP 200 OK. Status updated to 'completed', audit event logged.",
        "status": "PASS",
        "defect_id": "None",
        "evidence_link": "apps.scheduling.test_qa_suite.FullQASuiteTests.test_positive_full_journey"
    },
    {
        "id": "TC-POS-05",
        "feature": "Free Cancellation",
        "preconditions": "Booking confirmed, cancellation time > cancellation_hours window (24h).",
        "steps": "1. POST /api/bookings/{id}/cancel/ with optional reason.",
        "input": "booking_id=1, reason='Client schedule change'",
        "expected": "HTTP 200 OK. status='cancelled', is_late=False, slot released, reminders cancelled.",
        "actual": "HTTP 200 OK. status='cancelled', reminders marked cancelled.",
        "status": "PASS",
        "defect_id": "None",
        "evidence_link": "apps.scheduling.tests.SchedulingAndBookingEngineTests.test_cancellation_policy_enforcement"
    },
    {
        "id": "TC-NEG-01",
        "feature": "Negative Validation",
        "preconditions": "Customer authenticated.",
        "steps": "1. POST /api/bookings/ with start_at set in the past.",
        "input": "start_at='2026-10-01T10:00:00Z'",
        "expected": "HTTP 400 Bad Request. Validation error indicating appointment cannot be in the past.",
        "actual": "HTTP 400 Bad Request. Validation message: requires minimum notice.",
        "status": "PASS",
        "defect_id": "None",
        "evidence_link": "apps.scheduling.test_qa_suite.FullQASuiteTests.test_negative_past_date"
    },
    {
        "id": "TC-NEG-02",
        "feature": "Minimum Notice",
        "preconditions": "Provider min_notice_minutes=60.",
        "steps": "1. POST /api/bookings/ with start_at 20 minutes from now.",
        "input": "start_at=now + 20 minutes",
        "expected": "HTTP 400 Bad Request. Validation error indicating minimum 60 minutes notice required.",
        "actual": "HTTP 400 Bad Request with notice constraint error message.",
        "status": "PASS",
        "defect_id": "None",
        "evidence_link": "apps.scheduling.test_qa_suite.FullQASuiteTests.test_negative_insufficient_notice"
    },
    {
        "id": "TC-NEG-03",
        "feature": "Inactive Service",
        "preconditions": "Service has is_active=False.",
        "steps": "1. POST /api/bookings/ referencing inactive service_id.",
        "input": "service_id=2 (inactive)",
        "expected": "HTTP 400 Bad Request. 'The selected service is currently inactive.'",
        "actual": "HTTP 400 Bad Request with inactive service error.",
        "status": "PASS",
        "defect_id": "None",
        "evidence_link": "apps.scheduling.test_qa_suite.FullQASuiteTests.test_negative_inactive_service"
    },
    {
        "id": "TC-NEG-04",
        "feature": "Cross-Provider Service",
        "preconditions": "Service belongs to Provider B, request sent for Provider A.",
        "steps": "1. POST /api/bookings/ with provider_id=A and service_id=B.",
        "input": "provider_id=1, service_id=5 (Provider B)",
        "expected": "HTTP 400 Bad Request. 'The requested service does not belong to this provider.'",
        "actual": "HTTP 400 Bad Request. Validation rejected cross-provider mismatch.",
        "status": "PASS",
        "defect_id": "None",
        "evidence_link": "apps.scheduling.test_qa_suite.FullQASuiteTests.test_negative_wrong_service"
    },
    {
        "id": "TC-BND-01",
        "feature": "Double-Booking Prevention",
        "preconditions": "Booking exists for [10:00, 10:30).",
        "steps": "1. Attempt booking for overlapping interval [10:15, 10:45).",
        "input": "start_at='2026-11-05T10:15:00Z'",
        "expected": "HTTP 409 Conflict. Caught by PostgreSQL exclusion constraint no_overlap_active_bookings.",
        "actual": "HTTP 409 Conflict. Response: {'error': 'slot_conflict'}.",
        "status": "PASS",
        "defect_id": "None",
        "evidence_link": "apps.scheduling.test_qa_suite.FullQASuiteTests.test_boundary_back_to_back_bookings"
    },
    {
        "id": "TC-BND-02",
        "feature": "Contiguous Boundary",
        "preconditions": "Booking 1 ends at 10:30 [10:00, 10:30).",
        "steps": "1. Book adjacent slot starting exactly at 10:30 [10:30, 11:00).",
        "input": "start_at='2026-11-05T10:30:00Z'",
        "expected": "HTTP 201 Created. Half-open range '[)' allows exact boundary touch without collision.",
        "actual": "HTTP 201 Created. Both contiguous bookings coexist in database.",
        "status": "PASS",
        "defect_id": "None",
        "evidence_link": "apps.scheduling.test_qa_suite.FullQASuiteTests.test_boundary_back_to_back_bookings"
    },
    {
        "id": "TC-BND-03",
        "feature": "DST Transition",
        "preconditions": "America/New_York clocks fall back 1 hour on Nov 1, 2026.",
        "steps": "1. Generate slots across full 24h transition day.",
        "input": "date=2026-11-01, tz=America/New_York",
        "expected": "No crash or duplicate slot timestamps; monotonically increasing UTC slots.",
        "actual": "Slots generated smoothly with strict monotonic progression.",
        "status": "PASS",
        "defect_id": "None",
        "evidence_link": "apps.scheduling.test_qa_suite.FullQASuiteTests.test_boundary_dst_clock_shift"
    },
    {
        "id": "TC-SEC-01",
        "feature": "Cross-Tenant Privacy",
        "preconditions": "Booking A belongs to Customer A. Customer B authenticated.",
        "steps": "1. Customer B attempts GET /api/bookings/{id_A}/.",
        "input": "booking_id=1, user=Customer B",
        "expected": "HTTP 404 Not Found. Customer B cannot view Customer A's booking.",
        "actual": "HTTP 404 Not Found. Scoped queryset blocks access.",
        "status": "PASS",
        "defect_id": "None",
        "evidence_link": "apps.scheduling.test_security.SecurityAndReliabilityPhase4Tests.test_customer_cannot_view_another_customer_booking"
    },
    {
        "id": "TC-SEC-02",
        "feature": "Reschedule Hijack Defense",
        "preconditions": "Booking A belongs to Customer A.",
        "steps": "1. Customer B attempts PATCH /api/bookings/{id_A}/reschedule/.",
        "input": "booking_id=1, user=Customer B",
        "expected": "HTTP 403 Forbidden. IsBookingParticipant blocks unauthorized reschedule.",
        "actual": "HTTP 403 Forbidden. PermissionDenied raised.",
        "status": "PASS",
        "defect_id": "None",
        "evidence_link": "apps.scheduling.test_security.SecurityAndReliabilityPhase4Tests.test_customer_cannot_reschedule_another_customer_booking"
    },
    {
        "id": "TC-SEC-03",
        "feature": "Cancellation Hijack Defense",
        "preconditions": "Booking A belongs to Customer A.",
        "steps": "1. Customer B attempts POST /api/bookings/{id_A}/cancel/.",
        "input": "booking_id=1, user=Customer B",
        "expected": "HTTP 403 Forbidden. IsBookingParticipant blocks unauthorized cancellation.",
        "actual": "HTTP 403 Forbidden. PermissionDenied raised.",
        "status": "PASS",
        "defect_id": "None",
        "evidence_link": "apps.scheduling.test_security.SecurityAndReliabilityPhase4Tests.test_customer_cannot_cancel_another_customer_booking"
    },
    {
        "id": "TC-SEC-04",
        "feature": "Provider Data Isolation",
        "preconditions": "Provider A owns Service 1. Provider B authenticated.",
        "steps": "1. Provider B attempts GET/PUT/DELETE /api/providers/me/services/1/.",
        "input": "service_id=1, user=Provider B",
        "expected": "HTTP 404 Not Found. Provider B cannot access or mutate Provider A's service.",
        "actual": "HTTP 404 Not Found. Provider profile scoping protects all records.",
        "status": "PASS",
        "defect_id": "None",
        "evidence_link": "apps.scheduling.test_security.SecurityAndReliabilityPhase4Tests.test_provider_cannot_access_or_modify_another_provider_services"
    },
    {
        "id": "TC-SEC-05",
        "feature": "Role Separation",
        "preconditions": "Booking exists.",
        "steps": "1. Customer attempts POST /api/bookings/{id}/complete/.",
        "input": "booking_id=1, user=Customer",
        "expected": "HTTP 403 Forbidden. Only the assigned provider can complete the session.",
        "actual": "HTTP 403 Forbidden. IsBookingProvider permission strictly enforced.",
        "status": "PASS",
        "defect_id": "None",
        "evidence_link": "apps.scheduling.test_security.SecurityAndReliabilityPhase4Tests.test_customer_cannot_mark_booking_complete"
    },
    {
        "id": "TC-CONC-01",
        "feature": "10-Thread Concurrency Benchmark",
        "preconditions": "1 open slot. 10 distinct customer threads execute concurrent bookings.",
        "steps": "1. 10 simultaneous threads fire POST /api/bookings/ for identical start_at.",
        "input": "provider_id=1, service_id=1, start_at='2026-10-31T14:00:00Z', 10 parallel threads",
        "expected": "Exactly 1 thread receives HTTP 201 Created. 9 threads receive HTTP 409 Conflict. Database has exactly 1 active row.",
        "actual": "1x HTTP 201 Created, 9x HTTP 409 Conflict. Database active rows = 1. Zero double-bookings.",
        "status": "PASS",
        "defect_id": "None",
        "evidence_link": "DOCS/concurrency_test_evidence.txt"
    },
    {
        "id": "TC-REL-01",
        "feature": "Idempotency Replay",
        "preconditions": "Booking created with Idempotency-Key UUID.",
        "steps": "1. Resend exact same POST /api/bookings/ with identical key.",
        "input": "Idempotency-Key='UUID-12345'",
        "expected": "HTTP 200 OK returning existing booking. Zero duplicate database rows.",
        "actual": "HTTP 200 OK. Returned original booking ID, DB count remains 1.",
        "status": "PASS",
        "defect_id": "None",
        "evidence_link": "apps.scheduling.test_qa_suite.FullQASuiteTests.test_idempotent_replay"
    },
    {
        "id": "TC-REL-02",
        "feature": "Worker Crash & Retry Safety",
        "preconditions": "Reminder pending. Network failure simulated on dispatch.",
        "steps": "1. Worker run 1 fails -> check backoff. 2. Worker run 2 succeeds. 3. Worker run 3 checks duplicate protection.",
        "input": "process_due_reminders(simulate_failure=True), then retry",
        "expected": "Reminder transitions to backoff, then successfully sent, never duplicated on rerun.",
        "actual": "Backoff calculated as 2^attempts * 60s. Sent once, duplicate runs ignored.",
        "status": "PASS",
        "defect_id": "None",
        "evidence_link": "apps.scheduling.test_security.SecurityAndReliabilityPhase4Tests.test_worker_retry_safety_and_duplicate_send_protection"
    }
]


def generate_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Cases & Execution Matrix"
    ws.views.sheetView[0].showGridLines = True

    # Title Block
    ws.merge_cells("A1:J1")
    title_cell = ws["A1"]
    title_cell.value = "TimeSync SaaS - Comprehensive QA Test Execution Matrix"
    title_cell.font = Font(name="Calibri", size=16, bold=True, color="1E293B")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    ws.merge_cells("A2:J2")
    sub_cell = ws["A2"]
    sub_cell.value = "Zero Double-Booking Concurrency, Security, Boundary, and Full Lifecycle Tests"
    sub_cell.font = Font(name="Calibri", size=11, italic=True, color="64748B")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    headers = [
        "Test ID",
        "Feature",
        "Preconditions",
        "Steps",
        "Input",
        "Expected",
        "Actual",
        "Status",
        "Defect ID",
        "Evidence Link"
    ]

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid") # Deep Indigo
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.row_dimensions[4].height = 28
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    pass_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    pass_font = Font(name="Calibri", size=10, bold=True, color="065F46")

    data_font = Font(name="Calibri", size=10, color="0F172A")
    data_align = Alignment(vertical="center", wrap_text=True)
    center_align = Alignment(horizontal="center", vertical="center")

    for row_idx, tc in enumerate(test_cases_data, start=5):
        ws.row_dimensions[row_idx].height = 42
        row_values = [
            tc["id"],
            tc["feature"],
            tc["preconditions"],
            tc["steps"],
            tc["input"],
            tc["expected"],
            tc["actual"],
            tc["status"],
            tc["defect_id"],
            tc["evidence_link"]
        ]

        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.alignment = data_align
            cell.border = thin_border

            # Center align ID, Status, Defect ID
            if col_idx in (1, 8, 9):
                cell.alignment = center_align

            # Highlight PASS status
            if col_idx == 8 and val == "PASS":
                cell.fill = pass_fill
                cell.font = pass_font

    # Set column widths
    col_widths = {
        1: 14,  # Test ID
        2: 20,  # Feature
        3: 32,  # Preconditions
        4: 36,  # Steps
        5: 32,  # Input
        6: 34,  # Expected
        7: 34,  # Actual
        8: 12,  # Status
        9: 12,  # Defect ID
        10: 42  # Evidence Link
    }

    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Save outputs
    out_root = BASE_DIR / "05_test-cases.xlsx"
    out_docs = BASE_DIR / "DOCS" / "05_test-cases.xlsx"

    wb.save(out_root)
    wb.save(out_docs)
    print(f"Generated 05_test-cases.xlsx at:\n- {out_root}\n- {out_docs}")


if __name__ == "__main__":
    generate_excel()
